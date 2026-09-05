import os
import csv
import logging
from typing import List, Dict, Any, Optional
from pydantic import BaseModel
from app.core.config import settings

logger = logging.getLogger("jarvis.rag.parsers")

class DocumentElement(BaseModel):
    text: str
    page_number: Optional[int] = None
    section_name: Optional[str] = None
    sheet_name: Optional[str] = None
    slide_number: Optional[int] = None
    row_start: Optional[int] = None
    row_end: Optional[int] = None
    line_start: Optional[int] = None
    line_end: Optional[int] = None
    
    # Phase 9: OCR Metadata
    source_type: str = "NATIVE"
    confidence: Optional[float] = None
    bbox: Optional[Dict[str, float]] = None
    block_index: Optional[int] = None
    
    metadata: Dict[str, Any] = {}

class DocumentContent(BaseModel):
    document_id: str
    elements: List[DocumentElement]

class BaseDocumentParser:
    def can_parse(self, mime_type: str, extension: str) -> bool:
        raise NotImplementedError

    async def parse(self, file_path: str, document_id: str) -> DocumentContent:
        raise NotImplementedError

class TXTParser(BaseDocumentParser):
    def can_parse(self, mime_type: str, extension: str) -> bool:
        return extension.lower() in [".txt", ".md"] or "text/plain" in mime_type

    async def parse(self, file_path: str, document_id: str) -> DocumentContent:
        elements = []
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            lines = f.readlines()
            current_p = []
            start_line = 1
            for i, line in enumerate(lines):
                line = line.strip()
                if not line:
                    if current_p:
                        elements.append(DocumentElement(
                            text="\n".join(current_p),
                            line_start=start_line,
                            line_end=i
                        ))
                        current_p = []
                    start_line = i + 2
                else:
                    current_p.append(line)
            if current_p:
                elements.append(DocumentElement(
                    text="\n".join(current_p),
                    line_start=start_line,
                    line_end=len(lines)
                ))
        return DocumentContent(document_id=document_id, elements=elements)

class PDFParser(BaseDocumentParser):
    def can_parse(self, mime_type: str, extension: str) -> bool:
        return extension.lower() == ".pdf" or "application/pdf" in mime_type

    async def parse(self, file_path: str, document_id: str) -> DocumentContent:
        elements = []
        try:
            from pypdf import PdfReader
            from app.interfaces.ocr_provider import OCRProviderFactory
            from app.core_service.vision.pdf_renderer import PDFRenderer
            
            ocr_provider = OCRProviderFactory.get_provider()
            
            reader = PdfReader(file_path)
            for i, page in enumerate(reader.pages):
                page_number = i + 1
                text = page.extract_text()
                
                if text and text.strip():
                    elements.append(DocumentElement(
                        text=text.strip(),
                        page_number=page_number,
                        source_type="NATIVE"
                    ))
                else:
                    # Render and OCR this page
                    try:
                        image_data = PDFRenderer.render_page(file_path, page_number, settings.PDF_RENDER_DPI)
                        if image_data:
                            ocr_result = await ocr_provider.extract_text(image_data, "image/png")
                            if ocr_result.text and ocr_result.text.strip():
                                for b_idx, block in enumerate(ocr_result.blocks):
                                    elements.append(DocumentElement(
                                        text=block.text.strip(),
                                        page_number=page_number,
                                        source_type="OCR",
                                        confidence=block.confidence,
                                        bbox=block.bbox,
                                        block_index=b_idx
                                    ))
                    except Exception as ocr_err:
                        logger.error(f"OCR failed for page {page_number} of {document_id}: {str(ocr_err)}")
                        
        except ImportError:
            logger.error("pypdf missing")
            
        return DocumentContent(document_id=document_id, elements=elements)

class ImageParser(BaseDocumentParser):
    def can_parse(self, mime_type: str, extension: str) -> bool:
        return extension.lower() in [".png", ".jpg", ".jpeg", ".webp"] or mime_type.startswith("image/")

    async def parse(self, file_path: str, document_id: str) -> DocumentContent:
        elements = []
        try:
            with open(file_path, "rb") as f:
                image_data = f.read()
                
            from app.interfaces.ocr_provider import OCRProviderFactory
            ocr_provider = OCRProviderFactory.get_provider()
            
            # Simple mime_type detection
            ext = os.path.splitext(file_path)[1].lower()
            mime_type = "image/jpeg" if ext in [".jpg", ".jpeg"] else "image/png"
            if ext == ".webp": mime_type = "image/webp"
            
            ocr_result = await ocr_provider.extract_text(image_data, mime_type)
            
            if ocr_result.text and ocr_result.text.strip():
                for b_idx, block in enumerate(ocr_result.blocks):
                    elements.append(DocumentElement(
                        text=block.text.strip(),
                        page_number=1, # Images are treated as 1 page
                        source_type="OCR",
                        confidence=block.confidence,
                        bbox=block.bbox,
                        block_index=b_idx
                    ))
                    
        except Exception as e:
            logger.error(f"Image OCR failed for {document_id}: {str(e)}")
            
        return DocumentContent(document_id=document_id, elements=elements)

class DOCXParser(BaseDocumentParser):
    def can_parse(self, mime_type: str, extension: str) -> bool:
        return extension.lower() == ".docx"

    async def parse(self, file_path: str, document_id: str) -> DocumentContent:
        elements = []
        try:
            from docx import Document
            doc = Document(file_path)
            for p in doc.paragraphs:
                text = p.text.strip()
                if text:
                    elements.append(DocumentElement(text=text))
            for table in doc.tables:
                table_data = []
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    table_data.append(" | ".join(row_data))
                if table_data:
                    elements.append(DocumentElement(text="\n".join(table_data)))
        except ImportError:
            pass
        return DocumentContent(document_id=document_id, elements=elements)

class CSVParser(BaseDocumentParser):
    def can_parse(self, mime_type: str, extension: str) -> bool:
        return extension.lower() == ".csv"

    async def parse(self, file_path: str, document_id: str) -> DocumentContent:
        elements = []
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            reader = csv.reader(f)
            header = next(reader, None)
            row_idx = 1
            for row in reader:
                row_idx += 1
                text_parts = []
                for i, val in enumerate(row):
                    col_name = header[i] if header and i < len(header) else f"Col{i+1}"
                    text_parts.append(f"{col_name}: {val}")
                
                elements.append(DocumentElement(
                    text=", ".join(text_parts),
                    row_start=row_idx,
                    row_end=row_idx
                ))
        return DocumentContent(document_id=document_id, elements=elements)

class XLSXParser(BaseDocumentParser):
    def can_parse(self, mime_type: str, extension: str) -> bool:
        return extension.lower() == ".xlsx"

    async def parse(self, file_path: str, document_id: str) -> DocumentContent:
        elements = []
        try:
            from openpyxl import load_workbook
            wb = load_workbook(file_path, data_only=True, read_only=True)
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                    if any(cell is not None for cell in row):
                        row_text = " | ".join([str(c) if c is not None else "" for c in row])
                        elements.append(DocumentElement(
                            text=row_text,
                            sheet_name=sheet_name,
                            row_start=row_idx,
                            row_end=row_idx
                        ))
        except ImportError:
            pass
        return DocumentContent(document_id=document_id, elements=elements)

class PPTXParser(BaseDocumentParser):
    def can_parse(self, mime_type: str, extension: str) -> bool:
        return extension.lower() == ".pptx"

    async def parse(self, file_path: str, document_id: str) -> DocumentContent:
        elements = []
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            for i, slide in enumerate(prs.slides, start=1):
                slide_text = []
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        slide_text.append(shape.text.strip())
                if slide_text:
                    elements.append(DocumentElement(
                        text="\n".join(slide_text),
                        slide_number=i
                    ))
        except ImportError:
            pass
        return DocumentContent(document_id=document_id, elements=elements)

class DocumentParserRegistry:
    def __init__(self):
        self.parsers = [
            PDFParser(),
            ImageParser(),
            DOCXParser(),
            XLSXParser(),
            CSVParser(),
            PPTXParser(),
            TXTParser() # Fallback for text
        ]
        
    def get_parser(self, mime_type: str, extension: str) -> Optional[BaseDocumentParser]:
        for parser in self.parsers:
            if parser.can_parse(mime_type, extension):
                return parser
        return None

registry = DocumentParserRegistry()
